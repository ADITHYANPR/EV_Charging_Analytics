"""Build Complete Academic Case-Study Report in Word (DOCX) Format.

Title: Data-Driven Analytics for EV Charging Demand Forecasting and Anomaly Detection in Smart Grids
Course: MDI3006 Advanced Data Analytics
Complies fully with the 14-page official faculty specification.
"""

from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_PATH = Path("EV_Charging_Analytics_Case_Study_Draft.docx")
FIG_DIR = Path("outputs/figures/report")
PROJ_FIG_DIR = Path("outputs/figures")

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in twentieths of a point (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set cell background color."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="D3D3D3"):
    """Set clean subtle borders for table."""
    tblPr = table._tbl.tblPr
    borders_elm = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_elm)

def add_header_footer(doc):
    """Add clean running header and footer."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    
    # Normal Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("MDI3006 Advanced Data Analytics | EV Charging Demand Forecasting & Anomaly Detection")
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(120, 144, 156)
    
    # Normal Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Page ")
    frun.font.size = Pt(9)
    frun.font.color.rgb = RGBColor(100, 116, 139)
    # Add page number xml field
    f_p = fp._p
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    f_p.append(fldSimple)
    frun2 = fp.add_run(" | Comprehensive Academic Case Study Report")
    frun2.font.size = Pt(9)
    frun2.font.color.rgb = RGBColor(100, 116, 139)

def format_paragraph(p, space_before=4, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_heading_1(doc, text):
    h = doc.add_heading(text, level=1)
    format_paragraph(h, space_before=16, space_after=8)
    h.style.font.name = 'Arial'
    h.style.font.size = Pt(15)
    h.style.font.bold = True
    h.style.font.color.rgb = RGBColor(30, 58, 138)
    return h

def add_heading_2(doc, text):
    h = doc.add_heading(text, level=2)
    format_paragraph(h, space_before=12, space_after=6)
    h.style.font.name = 'Arial'
    h.style.font.size = Pt(12.5)
    h.style.font.bold = True
    h.style.font.color.rgb = RGBColor(14, 116, 144)
    return h

def add_heading_3(doc, text):
    h = doc.add_heading(text, level=3)
    format_paragraph(h, space_before=8, space_after=4)
    h.style.font.name = 'Arial'
    h.style.font.size = Pt(11)
    h.style.font.bold = True
    h.style.font.color.rgb = RGBColor(51, 65, 85)
    return h

def add_figure_with_caption(doc, img_path, fig_num, caption_title, caption_desc, width_in=6.2):
    """Embed an image with a publication-grade caption and explanatory text."""
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_img, space_before=10, space_after=4)
    if Path(img_path).exists():
        p_img.add_run().add_picture(str(img_path), width=Inches(width_in))
    else:
        p_img.add_run(f"[FIGURE FILE NOT FOUND: {img_path}]").font.color.rgb = RGBColor(220, 38, 38)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_cap, space_before=2, space_after=10)
    
    r_num = p_cap.add_run(f"Figure {fig_num}: {caption_title}. ")
    r_num.font.bold = True
    r_num.font.size = Pt(9.5)
    r_num.font.color.rgb = RGBColor(30, 58, 138)
    
    r_desc = p_cap.add_run(caption_desc)
    r_desc.font.size = Pt(9)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(71, 85, 105)

def style_academic_table(table, col_widths, hdr_bg="1E3A8A", alt_bg="F8FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "CBD5E1")
    
    # Header Row
    for idx, cell in enumerate(table.rows[0].cells):
        set_cell_shading(cell, hdr_bg)
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(255, 255, 255)
                
    # Data Rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8.0)
                    r.font.color.rgb = RGBColor(30, 41, 59)
                    
    # Set explicit widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

def add_table_caption(doc, tbl_num, caption_title):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=14, space_after=4)
    r1 = p.add_run(f"Table {tbl_num}: ")
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = RGBColor(30, 58, 138)
    r2 = p.add_run(caption_title)
    r2.font.bold = True
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(30, 41, 59)
    return p

print("Loaded document builder utilities successfully.")
